"""Saved OOXML, cached native pivot, layout, and literal-text regression checks."""

import io
from collections import defaultdict
from dataclasses import replace
from datetime import datetime
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from scripts.generate_report_samples import sample_snapshot
from services.report_model import ReportSnapshot, build_report_data
from services.report_xlsx import build_report_xlsx, finalize_layout, money_cell
from services.report_formatting import indian_number, indian_format


def exported(source=None):
    data = build_report_data(source or sample_snapshot())
    payload = build_report_xlsx(data)
    return data, load_workbook(io.BytesIO(payload)), payload


def test_native_pivot_cache_and_saved_totals_match_allocation_source():
    data, workbook, payload = exported()
    assert workbook.sheetnames == ['Summary', 'Members & Families', 'Transactions', 'Payments']
    transactions = workbook['Transactions']
    assert transactions.tables['ReportAllocations'].ref == f'A4:T{4 + len(data.allocations)}'
    pivot = workbook['Summary']._pivots[0]
    assert pivot.cache.cacheSource.worksheetSource.name == 'ReportAllocations'
    assert pivot.cache.saveData is True and pivot.cache.refreshOnLoad is False
    assert pivot.cache.recordCount == len(data.allocations)
    assert len(pivot.cache.records.r) == len(data.allocations)
    assert [field.x for field in pivot.rowFields] == [6, 7]
    assert [field.x for field in pivot.colFields] == [2]
    assert [field.fld for field in pivot.pageFields] == [1, 4]
    assert [(field.fld, field.subtotal) for field in pivot.dataFields] == [(10, 'sum')]
    assert pivot.rowGrandTotals and pivot.colGrandTotals and pivot.enableDrill
    assert pivot.pivotFields[6].defaultSubtotal
    assert not pivot.pivotFields[7].defaultSubtotal
    assert len(pivot.cache.cacheFields) == 20
    assert not {'Account ID', 'Person ID'} & {f.name for f in pivot.cache.cacheFields}

    # Independently group the serialised cache records using their shared-item indices.
    totals = defaultdict(int)
    for record in pivot.cache.records.r:
        items = record._fields
        category = pivot.cache.cacheFields[2].sharedItems._fields[items[2].v].v
        amount = pivot.cache.cacheFields[10].sharedItems._fields[items[10].v].v
        totals[category] += amount
    assert dict(totals) == data.category_totals
    sheet = workbook['Summary']
    grand = next(row for row in sheet if row[0].value == 'Grand Total')
    assert [cell.value for cell in grand[2:3 + len(data.categories)]] == [*[data.category_totals[c] for c in data.categories], data.net]
    assert len(pivot.rowItems) == len({(a.account_id, a.person_id) for a in data.allocations}) + len(data.accounts) + 1
    with ZipFile(io.BytesIO(payload)) as archive:
        names = archive.namelist()
        assert 'xl/pivotTables/pivotTable1.xml' in names
        assert 'xl/pivotCache/pivotCacheRecords1.xml' in names
        assert len([n for n in names if n.startswith('xl/charts/chart') and n.endswith('.xml')]) == 1
        # Every produced XML part remains well formed.
        for name in names:
            if name.endswith('.xml') or name.endswith('.rels'):
                ET.fromstring(archive.read(name))


def test_new_export_includes_new_rows_in_named_source_and_cache():
    source = sample_snapshot()
    extra = {**source.expenses[0], 'id': 'new-record', 'amount': 24}
    data, workbook, _ = exported(replace(source, expenses=(*source.expenses, extra)))
    assert workbook['Summary']._pivots[0].cache.recordCount == len(data.allocations)
    assert any(row[11] == 'new-record' for row in workbook['Transactions'].values)
    assert len(data.allocations) > len(build_report_data(source).allocations)


def test_transactions_have_complete_metadata_on_participating_rows_and_sortable_dates():
    data, workbook, _ = exported()
    rows = list(workbook['Transactions'].values)[4:4 + len(data.allocations)]
    for row, allocation in zip(rows, data.allocations):
        assert row[0:11] == (allocation.number, allocation.occurred_at, allocation.category, allocation.description,
                             allocation.entry_type, allocation.payer, allocation.account, allocation.person,
                             allocation.split_method, allocation.expense_amount, allocation.share)
        assert row[11] == allocation.expense_id
        assert isinstance(row[1], datetime)
    assert sum(row[10] for row in rows) == data.net
    exact = [r for r in rows if r[11] == 'e04']
    assert sorted(r[10] for r in exact) == [0, 1000, 1000, 1200]
    fx = next(r for r in rows if r[11] == 'e07')
    assert fx[12:15] == (100, 'USD', '82.000000000000')
    assert fx[15] == datetime(2026, 10, 2)
    assert fx[17] == 'manual'
    assert workbook['Transactions']['M5'].number_format == '#,##0.###############'
    assert workbook['Transactions']['P5'].number_format == 'dd mmm yyyy'
    assert workbook['Transactions']['T5'].number_format == 'dd mmm yyyy'
    assert workbook['Transactions'].column_dimensions['L'].hidden
    assert workbook['Transactions'].freeze_panes == 'D5'


def test_literal_user_text_cannot_become_excel_formulas_in_any_sheet():
    source = sample_snapshot()
    source.trip['name'] = '=1+2'
    source.trip['members'][0]['name'] = '=HYPERLINK("https://invalid.example")'
    expenses = tuple({**e, 'description': '=2+3', 'category': '=4+5'} for e in source.expenses)
    _, workbook, _ = exported(replace(source, expenses=expenses))
    formula_cells = [(s.title, c.coordinate) for s in workbook for row in s for c in row if c.data_type == 'f']
    assert formula_cells == []
    assert workbook['Summary']['A1'].value == '=1+2'
    assert any(row[2] == '=4+5' and row[3] == '=2+3' for row in workbook['Transactions'].values)


def test_chart_uses_whole_trip_cache_and_receiving_paying_colours_are_semantic():
    data, workbook, _ = exported()
    series = workbook['Summary']._charts[0].series[0]
    assert [v.v for v in series.cat.strRef.strCache.pt] == list(data.categories)
    assert [v.v for v in series.val.numRef.numCache.pt] == [data.category_totals[c] for c in data.categories]
    assert len(series.data_points) == 1
    assert series.data_points[0].spPr.solidFill.srgbClr == '166534'
    rules = [rule for values in workbook['Members & Families'].conditional_formatting._cf_rules.values() for rule in values]
    assert {r.operator: r.dxf.font.color.rgb[-6:] for r in rules if r.type == 'cellIs' and r.dxf.fill} == {
        'greaterThan': '166534', 'lessThan': '991B1B', 'equal': '475569',
    }
    # Filtering a six-digit total down below one lakh must not display leading zeros.
    pivot_rules = [r for values in workbook['Summary'].conditional_formatting._cf_rules.values() for r in values]
    assert any(r.type == 'expression' and '<100000)' in r.formula[0] and r.dxf.numFmt.formatCode == '#,##0;(#,##0);0' for r in pivot_rules)


def test_account_totals_exclude_children_and_payments_keep_gross_source_totals():
    data, workbook, _ = exported()
    sheet = workbook['Members & Families']
    total = next(row for row in sheet.values if row[0] == 'Total')
    assert total[1:] == (data.net, data.net, 0, 12000, 12000, 0, 0, None)
    family = next(row for row in sheet.values if row[0] == 'Patel family')
    assert family[1] == 82905 and family[4:6] == (2000, 9400)
    history = list(workbook['Payments'].values)
    assert next(r[2] for r in history if r[0] == 'History total') == 12000
    assert next(r[2] for r in history if r[0] == 'Recorded payments total') == 7000
    assert next(r[2] for r in history if r[0] == 'Legacy settlements total') == 5000
    assert sum(r[2] for r in history if r[3] == 'Suggested') == data.outstanding
    assert sum(r[7] == 'payment-4' or r[7] == 'payment-5' for r in history) == 2
    assert not any(r[7] == 'legacy-pending' for r in history)


def test_numeric_indian_masks_and_long_labels_fit_without_converting_money_to_text():
    assert indian_number(272628) == '2,72,628'
    assert indian_number(-9876543210123) == '(98,76,54,32,10,123)'
    assert indian_format(272628) == r'0\,00\,000;(0\,00\,000);0'
    _, workbook, _ = exported()
    sheet = workbook['Summary']
    assert isinstance(sheet['G8'].value, int)
    assert r'\,' in sheet['G8'].number_format
    sheet['A70'] = 'Participant with a deliberately long display name ' * 4
    sheet['A70'].alignment = Alignment(wrap_text=True)
    cell = sheet.cell(70, 2, -9876543210123)
    money_cell(cell)
    finalize_layout(sheet, (2,))
    assert sheet.row_dimensions[70].height > 30
    assert sheet.column_dimensions['B'].width >= len(indian_number(cell.value)) + 3


def test_empty_report_has_four_tabs_no_dummy_pivot_source_and_no_chart():
    source = ReportSnapshot({'id': 'empty', 'name': 'Empty', 'members': [], 'currency': 'INR'})
    _, workbook, _ = exported(source)
    assert workbook.sheetnames == ['Summary', 'Members & Families', 'Transactions', 'Payments']
    assert not workbook['Summary']._pivots and not workbook['Summary']._charts
    assert not workbook['Transactions'].tables
    assert any(row[0] == 'No expenses recorded' for row in workbook['Summary'].values)


def test_missing_date_matches_saved_source_and_native_cache():
    source = sample_snapshot()
    expenses = tuple({**e, 'date': 'unavailable'} for e in source.expenses)
    _, workbook, _ = exported(replace(source, expenses=expenses))
    pivot = workbook['Summary']._pivots[0]
    assert [item.v for item in pivot.cache.cacheFields[1].sharedItems._fields] == ['unavailable']
    assert workbook['Transactions']['B5'].value == 'unavailable'


def test_wide_category_pivot_does_not_overlap_or_hide_chart_source_columns():
    source = sample_snapshot()
    expenses = tuple({**source.expenses[0], 'id': str(i), 'category': f'Category {i:02d}', 'amount': 20}
                     for i in range(30))
    data, workbook, _ = exported(ReportSnapshot(source.trip, expenses))
    summary = workbook['Summary']
    assert not summary.column_dimensions['AA'].hidden
    assert summary.column_dimensions['AJ'].hidden
    assert summary._charts[0].series[0].cat.strRef.f.startswith("'Summary'!$AJ$")
    grand = next(row for row in summary.values if row[0] == 'Grand Total')
    assert grand[32] == data.net


def test_an_unrecorded_time_is_displayed_as_a_date_without_invented_midnight():
    source = sample_snapshot()
    expenses = tuple({**e, 'time': None} for e in source.expenses)
    _, workbook, _ = exported(replace(source, expenses=expenses))
    assert workbook['Transactions']['B5'].number_format == 'dd mmm yyyy'


def test_report_metadata_is_compact_and_reconciliation_stays_in_balance_table():
    data, workbook, _ = exported()
    values = [c.value for row in workbook['Summary'] for c in row if isinstance(c.value, str)]
    assert not {'Members', 'Share code', 'Currency', 'Exported at (IST)'} & set(values)
    assert data.metadata.share_code not in values
    assert not any('Whole-history snapshot' in value for value in values)
    assert workbook['Summary']['A5'].value == 'Budget (₹)'
    assert workbook['Summary']['G8'].value == data.net
    assert 'Adjustments' not in workbook.sheetnames
    balances = workbook['Members & Families']
    assert balances['G4'].value == 'Adjustments (₹)'
    assert next(row[6] for row in balances.values if row[0] == 'Patel family') == 7
    notes = [c.value for row in balances for c in row if isinstance(c.value, str)]
    assert all(any(a.explanation in note for note in notes) for a in data.adjustments)
    assert not any(c.hyperlink for sheet in workbook for row in sheet for c in row)


@pytest.mark.parametrize('currency,symbol', [('AED', 'AED'), ('EUR', '€')])
def test_each_financial_section_uses_the_trip_currency(currency, symbol):
    source = sample_snapshot()
    data, workbook, _ = exported(replace(source, trip={**source.trip, 'currency': currency}))
    summary = workbook['Summary']
    assert summary['A5'].value == f'Budget ({symbol})'
    assert [summary.cell(7, col).value for col in (1, 4, 7, 10)] == [
        f'{label} ({symbol})' for label in ['Gross spending', 'Money back', 'Net spending', 'Still to settle']]
    assert [summary.cell(11, col).value for col in (2, 3, 4)] == [
        f'{label} ({symbol})' for label in ['Gross funded', 'Money back', 'Net funded']]
    pivot = summary._pivots[0]
    assert pivot.dataCaption == f'Net expense share ({symbol})'
    assert pivot.dataFields[0].name == f'Expense share ({symbol})'
    assert [c.value for c in workbook['Members & Families'][4][1:-1]] == [
        f'{label} ({symbol})' for label in ['Expense share', 'Net funded', 'Before settlements',
                                         'Settlements paid', 'Settlements received', 'Adjustments', 'Balance now']]
    assert workbook['Transactions']['J4'].value == f'Expense total ({symbol})'
    assert workbook['Transactions']['K4'].value == f'Expense share ({symbol})'
    assert workbook['Payments']['C5'].value == f'Amount ({symbol})'
    assert sum(a.share for a in data.allocations) == summary['G8'].value
    assert not any('₹' in str(c.value) or 'INR' in str(c.value)
                   for sheet in workbook for row in sheet for c in row)
