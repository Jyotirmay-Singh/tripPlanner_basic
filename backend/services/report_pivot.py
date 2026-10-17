"""Native OOXML PivotTable with complete cache data and pre-rendered worksheet results.

    The adapter uses openpyxl's pivot serialization types and does not require Excel on the server.
    Its source is the generated named allocation table, never a bundled user workbook/template.
"""

from collections import defaultdict
from datetime import date, datetime

from openpyxl.pivot.cache import CacheDefinition, CacheSource, WorksheetSource, CacheField, SharedItems
from openpyxl.pivot.fields import Text, Number, DateTimeField, Missing, Index
from openpyxl.pivot.record import Record, RecordList
from openpyxl.pivot.table import (
    TableDefinition, Location, PivotField, FieldItem, RowColField, RowColItem,
    DataField, PageField, PivotTableStyle,
)
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.table import TableStyle, TableStyleElement

from services.report_formatting import BRAND, NEUTRAL, currency_symbol


def _text(sheet, row, column, value):
    cell = sheet.cell(row, column, str(value))
    cell.data_type = "s"
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    return cell


def _forest_style(workbook):
    elements = []
    for kind, fill, color, bold in (
        ("wholeTable", "FFFFFF", "121A18", False),
        ("headerRow", BRAND, "FFFFFF", True),
        ("firstSubtotalRow", "EAF0EE", BRAND, True),
        ("secondSubtotalRow", "EAF0EE", BRAND, True),
        ("totalRow", "DDE9E5", BRAND, True),
        ("firstRowStripe", "F4F7F6", "121A18", False),
        ("secondRowStripe", "FFFFFF", "121A18", False),
        ("pageFieldLabels", "FFFFFF", BRAND, True),
        ("pageFieldValues", "FFFFFF", NEUTRAL, False),
    ):
        style = DifferentialStyle(font=Font(name="Calibri", size=11, color=color, bold=bold),
                                  fill=PatternFill("solid", fgColor=fill))
        elements.append(TableStyleElement(type=kind, dxfId=workbook._differential_styles.add(style),
                                          size=1 if "Stripe" in kind else None))
    workbook._table_styles.tableStyle.append(TableStyle(name="ReportPivotForest", pivot=True, table=False,
                                                        tableStyleElement=elements))
    return "ReportPivotForest"


def add_allocation_pivot(workbook, sheet, source_headers, source_rows, data, start_row, style_money, style_header):
    """Write a fully cached account/person/category pivot; return its occupied bottom row."""
    if not source_rows:
        sheet.cell(start_row, 1, "No expenses recorded")
        return start_row
    account_col, person_col, category_col, amount_col = 6, 7, 2, 10
    date_col, type_col = 1, 4
    shared, indexes, cache_fields = [], [], []
    for column, header in enumerate(source_headers):
        values, lookup = [], {}
        for row in source_rows:
            value = row[column]
            key = (type(value).__name__, value)
            if key not in lookup:
                lookup[key] = len(values)
                values.append(value)
        fields = []
        for value in values:
            if value is None:
                fields.append(Missing())
            elif isinstance(value, (date, datetime)):
                fields.append(DateTimeField(v=datetime.combine(value, datetime.min.time()) if not isinstance(value, datetime) else value))
            elif isinstance(value, (int, float)):
                fields.append(Number(v=value))
            else:
                fields.append(Text(v=str(value)))
        contains_date = any(isinstance(value, (date, datetime)) for value in values)
        contains_number = any(isinstance(value, (int, float)) for value in values)
        contains_text = any(isinstance(value, str) for value in values)
        cache_fields.append(CacheField(name=header, sharedItems=SharedItems(
            _fields=fields, containsDate=contains_date, containsNonDate=any(not isinstance(v, (date, datetime)) for v in values),
            containsNumber=contains_number,
            containsInteger=contains_number and all(not isinstance(v, (int, float)) or float(v).is_integer() for v in values),
            containsString=contains_text, containsBlank=None in values,
            containsMixedTypes=sum((contains_date, contains_number, contains_text)) > 1,
            containsSemiMixedTypes=None in values,
        )))
        shared.append(values)
        indexes.append(lookup)
    records = RecordList(r=[Record(_fields=[Index(v=indexes[col][(type(value).__name__, value)])
                                           for col, value in enumerate(row)]) for row in source_rows])
    cache = CacheDefinition(
        cacheSource=CacheSource(type="worksheet", worksheetSource=WorksheetSource(name="ReportAllocations")),
        cacheFields=cache_fields, recordCount=len(source_rows), saveData=True, refreshOnLoad=False,
        enableRefresh=True, createdVersion=6, refreshedVersion=6, minRefreshableVersion=3,
        missingItemsLimit=0,
    )
    cache.records = records
    pivot_fields = []
    for col, values in enumerate(shared):
        axis = "axisRow" if col in (account_col, person_col) else "axisCol" if col == category_col else "axisPage" if col in (date_col, type_col) else None
        items = [FieldItem(x=index) for index in range(len(values))] if axis else []
        if col == account_col:
            items.append(FieldItem(t="default"))
        pivot_fields.append(PivotField(axis=axis, dataField=col == amount_col, items=items,
                                       defaultSubtotal=col == account_col, subtotalTop=False,
                                       showAll=False, compact=False, outline=False))

    totals = defaultdict(lambda: defaultdict(int))
    account_ids = []
    person_order = defaultdict(list)
    for allocation in data.allocations:
        if allocation.account_id not in account_ids:
            account_ids.append(allocation.account_id)
        key = (allocation.account_id, allocation.person_id)
        if key not in person_order[allocation.account_id]:
            person_order[allocation.account_id].append(key)
        totals[key][allocation.category] += allocation.share
    labels = {(row.account_id, row.person_id): (row.account, row.person) for row in data.allocations}
    account_names = {row.account_id: row.name for row in data.accounts}
    categories = list(data.categories)
    last_col = 3 + len(categories)
    # Report filters occupy their own rows outside the pivot's display region.
    sheet.cell(start_row - 3, 1, source_headers[date_col])
    sheet.cell(start_row - 3, 2, "(All)")
    sheet.cell(start_row - 2, 1, source_headers[type_col])
    sheet.cell(start_row - 2, 2, "(All)")
    measure_caption = f"Net expense share ({currency_symbol(data.currency)})"
    sheet.cell(start_row, 1, measure_caption)
    sheet.cell(start_row, 3, "Category")
    captions = ["Family / account", "Person", *categories, f"Total ({currency_symbol(data.currency)})"]
    for col, caption in enumerate(captions, 1):
        _text(sheet, start_row + 1, col, caption)
    style_header(sheet, start_row + 1, last_col)
    row_items, row_no = [], start_row + 2
    for account_id in account_ids:
        account_index = shared[account_col].index(account_names[account_id])
        subtotal = defaultdict(int)
        for key in person_order[account_id]:
            account, person = labels[key]
            person_index = shared[person_col].index(person)
            values = [totals[key][category] for category in categories]
            _text(sheet, row_no, 1, account)
            _text(sheet, row_no, 2, person)
            for col, value in enumerate([*values, sum(values)], 3):
                style_money(sheet.cell(row_no, col, value))
            row_items.append(RowColItem(x=[Index(v=account_index), Index(v=person_index)]))
            for category in categories:
                subtotal[category] += totals[key][category]
            row_no += 1
        _text(sheet, row_no, 1, account_names[account_id] + " Total")
        for col, value in enumerate([*[subtotal[c] for c in categories], sum(subtotal.values())], 3):
            style_money(sheet.cell(row_no, col, value))
        row_items.append(RowColItem(t="default", x=[Index(v=account_index)]))
        row_no += 1
    sheet.cell(row_no, 1, "Grand Total")
    for col, value in enumerate([*[data.category_totals[c] for c in categories], data.net], 3):
        style_money(sheet.cell(row_no, col, value))
    row_items.append(RowColItem(t="grand", x=[Index(v=0)]))
    column_items = [RowColItem(x=[Index(v=shared[category_col].index(category))]) for category in categories]
    column_items.append(RowColItem(t="grand", x=[Index(v=0)]))
    numeric_format_id = workbook._number_formats.add("#,##0;(#,##0);0") + 164
    pivot = TableDefinition(
        name="ExpenseShareByPersonCategory", cacheId=1, dataCaption=measure_caption,
        grandTotalCaption="Grand Total", rowHeaderCaption="Family / account", colHeaderCaption="Category",
        location=Location(ref=f"A{start_row}:{get_column_letter(last_col)}{row_no}",
                          firstHeaderRow=1, firstDataRow=2, firstDataCol=2,
                          rowPageCount=2, colPageCount=1),
        pivotFields=pivot_fields, rowFields=[RowColField(x=account_col), RowColField(x=person_col)],
        rowItems=row_items, colFields=[RowColField(x=category_col)], colItems=column_items,
        pageFields=[PageField(fld=date_col, hier=-1), PageField(fld=type_col, hier=-1)],
        dataFields=[DataField(name=source_headers[amount_col], fld=amount_col, numFmtId=numeric_format_id)],
        rowGrandTotals=True, colGrandTotals=True, compact=False, compactData=False, outline=False,
        outlineData=False, gridDropZones=False, showHeaders=True, showDrill=True, enableDrill=True,
        preserveFormatting=True, useAutoFormatting=True, applyNumberFormats=True,
        createdVersion=6, updatedVersion=6, minRefreshableVersion=3,
        pivotTableStyleInfo=PivotTableStyle(name=_forest_style(workbook), showRowHeaders=True,
                                          showColHeaders=True, showRowStripes=True,
                                          showColStripes=False, showLastColumn=True),
    )
    pivot.cache = cache
    sheet.add_pivot(pivot)
    return row_no
