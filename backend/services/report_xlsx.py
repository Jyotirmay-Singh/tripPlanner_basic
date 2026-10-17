"""Professional snapshot workbook, rendered from the shared reconciled report model."""

import io
from collections import defaultdict
from datetime import datetime
from math import ceil

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabel, DataLabelList
from openpyxl.chart.data_source import AxDataSource, StrRef, StrData, StrVal, NumData, NumVal
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule, Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from services.report_formatting import BRAND, RECEIVE, PAY, NEUTRAL, currency_symbol, indian_format, indian_number
from services.report_model import ReportData, routing_label
from services.report_pivot import add_allocation_pivot


def text_cell(sheet, row, col, value):
    cell = sheet.cell(row, col, value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def money_cell(cell):
    cell.number_format = indian_format(int(cell.value or 0))
    cell.alignment = Alignment(horizontal="right", vertical="center")


def header_row(sheet, row, columns):
    for col in range(1, columns + 1):
        cell = sheet.cell(row, col)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BRAND)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 34


def _row(sheet, row, values):
    for col, value in enumerate(values, 1):
        cell = text_cell(sheet, row, col, value)
        if isinstance(value, (int, float)):
            money_cell(cell)
        elif isinstance(value, datetime):
            cell.number_format = 'dd mmm yyyy, h:mm AM/PM'
        cell.alignment = Alignment(vertical="center", horizontal="right" if isinstance(value, (int, float)) else "left", wrap_text=True)


def _total(sheet, row, columns):
    for col in range(1, columns + 1):
        cell = sheet.cell(row, col)
        cell.font = Font(name="Calibri", size=11, bold=True, color=BRAND)
        cell.fill = PatternFill("solid", fgColor="EAF0EE")
        cell.border = Border(top=Side(style="thin", color=BRAND))


def _title(sheet, title, columns):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    text_cell(sheet, 1, 1, title).font = Font(name="Calibri", size=18, bold=True, color=BRAND)
    sheet.row_dimensions[1].height = 28


def _note(sheet, row, message, columns=10):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
    cell = text_cell(sheet, row, 1, message)
    cell.font = Font(name="Calibri", size=10, color=NEUTRAL)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row].height = 30


def _widths(sheet, widths):
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width


def _table(sheet, name, header, last, columns):
    if last <= header:
        return
    table = Table(displayName=name, ref=f"A{header}:{get_column_letter(columns)}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
    sheet.add_table(table)


def _number_rules(sheet, area, anchor):
    # Filters can change pivot sums. Select a numeric Indian mask again for the new magnitude.
    style = DifferentialStyle(numFmt=NumberFormat(numFmtId=0, formatCode=indian_format(0)))
    sheet.conditional_formatting.add(area, Rule(type="expression", dxf=style,
        formula=[f"AND(ISNUMBER({anchor}),ABS({anchor})<100000)"]))
    for digits in range(6, 18, 2):
        low, high = 10 ** (digits - 1), 10 ** (digits + 1)
        style = DifferentialStyle(numFmt=NumberFormat(numFmtId=0, formatCode=indian_format(low)))
        sheet.conditional_formatting.add(area, Rule(type="expression", dxf=style,
            formula=[f"AND(ISNUMBER({anchor}),ABS({anchor})>={low},ABS({anchor})<{high})"]))


def _balance_rules(sheet, area):
    for operator, value, color, fill in (("greaterThan", 0, RECEIVE, "F0FDF4"),
                                         ("lessThan", 0, PAY, "FEF2F2"),
                                         ("equal", 0, NEUTRAL, "F1F5F9")):
        sheet.conditional_formatting.add(area, CellIsRule(operator=operator, formula=[str(value)],
            font=Font(color=color), fill=PatternFill("solid", fgColor=fill)))


def finalize_layout(sheet, money_columns=()):
    for row in sheet:
        height = sheet.row_dimensions[row[0].row].height or 18
        for cell in row:
            if cell.__class__.__name__ == "MergedCell":
                continue
            if cell.font.name is None:
                cell.font = Font(name="Calibri", size=11, color="121A18")
            if cell.column in money_columns and isinstance(cell.value, (int, float)):
                width = len(indian_number(cell.value)) + 3
                dim = sheet.column_dimensions[cell.column_letter]
                dim.width = max(dim.width or 13, width)
            if isinstance(cell.value, str) and cell.alignment.wrap_text:
                merged = any(cell.coordinate in span for span in sheet.merged_cells.ranges)
                if not merged:
                    width = max(8, int(sheet.column_dimensions[cell.column_letter].width or 13) - 2)
                    height = max(height, 15 * max(ceil(len(line) / width) for line in cell.value.splitlines() or [""]))
        sheet.row_dimensions[row[0].row].height = max(18, height)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.oddFooter.center.text = "Page &P of &N"


def _transactions(workbook, data):
    sheet = workbook.create_sheet("Transactions")
    symbol = currency_symbol(data.currency)
    headers = ["Expense No", "Date (IST)", "Category", "Description", "Entry type", "Paid by",
               "Family / account", "Person", "Split method", f"Expense total ({symbol})", f"Expense share ({symbol})",
               "Expense ID", "Original amount", "Original currency",
               "Stored exchange rate", "Effective rate date", "Rate provider", "Rate mode", "Original exact allocations", "Requested rate date"]
    _title(sheet, "Transactions", 11)
    _note(sheet, 2, "One row per participating person. Expense total is context only; sum Expense share. FX evidence is grouped to the right.", 11)
    _row(sheet, 4, headers)
    header_row(sheet, 4, len(headers))
    source_rows = []
    for number, allocation in enumerate(data.allocations, 5):
        row = [allocation.number, allocation.occurred_at or allocation.date_source or "Not recorded", allocation.category, allocation.description,
               allocation.entry_type, allocation.payer, allocation.account, allocation.person,
               allocation.split_method, allocation.expense_amount, allocation.share,
               allocation.expense_id,
               float(allocation.original_amount) if allocation.original_amount is not None else None, allocation.original_currency, allocation.exchange_rate,
               datetime.combine(allocation.rate_date, datetime.min.time()) if allocation.rate_date else None,
               allocation.rate_provider, allocation.rate_mode, allocation.original_exact,
               datetime.combine(allocation.rate_requested_date, datetime.min.time()) if allocation.rate_requested_date else None]
        source_rows.append(row)
        _row(sheet, number, row)
        if allocation.occurred_at and not allocation.time_recorded:
            sheet.cell(number, 2).number_format = 'dd mmm yyyy'
        sheet.cell(number, 1).number_format = "0"
        sheet.cell(number, 13).number_format = '#,##0.###############'
        sheet.cell(number, 16).number_format = "dd mmm yyyy"
        sheet.cell(number, 20).number_format = "dd mmm yyyy"
        for col in (10, 11):
            money_cell(sheet.cell(number, col))
    last = 4 + len(source_rows)
    _table(sheet, "ReportAllocations", 4, last, len(headers))
    if source_rows:
        _row(sheet, last + 2, ["Net expense shares", None, None, None, None, None, None, None, None, None, data.net])
        _total(sheet, last + 2, 11)
        sheet.conditional_formatting.add(f"K5:K{last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color=RECEIVE)))
        _number_rules(sheet, f"J5:K{last + 2}", "J5")
    else:
        text_cell(sheet, 5, 1, "No expenses recorded")
    _widths(sheet, [11, 24, 22, 38, 16, 23, 25, 24, 16, 19, 19, 25, 20, 18, 28, 20, 22, 18, 45, 20])
    sheet.column_dimensions.group("L", "T", hidden=True)
    sheet.freeze_panes = "D5"
    sheet.print_title_rows = "1:4"
    sheet.print_area = f"A1:K{max(last + 2, 5)}"
    return headers, source_rows


def _summary(workbook, data, headers, source_rows):
    sheet = workbook["Summary"]
    meta = data.metadata
    symbol = currency_symbol(data.currency)
    _widths(sheet, [25, 22, *([13] * max(10, len(data.categories) + 1))])
    _title(sheet, meta.trip_name, 12)
    metadata = [("Trip dates", meta.trip_dates), ("Included expense dates", meta.expense_dates),
                ("Included settlement dates", meta.settlement_dates),
                (f"Budget ({symbol})", meta.budget if meta.budget is not None else "Not set")]
    for row, (label, value) in enumerate(metadata, 2):
        text_cell(sheet, row, 1, label).font = Font(name="Calibri", size=11, bold=True, color=BRAND)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        cell = text_cell(sheet, row, 2, value)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if isinstance(value, datetime):
            cell.number_format = 'dd mmm yyyy, h:mm AM/PM'
        elif isinstance(value, (int, float)):
            money_cell(cell)
            cell.alignment = Alignment(horizontal="center")
    for col, (label, value) in zip((1, 4, 7, 10), (("Gross spending", data.gross), ("Money back", data.money_back),
                                                ("Net spending", data.net), ("Still to settle", data.outstanding))):
        for row in (7, 8):
            sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
        text_cell(sheet, 7, col, f"{label} ({symbol})").font = Font(name="Calibri", size=11, bold=True, color=BRAND)
        money_cell(sheet.cell(8, col, value))
        sheet.cell(8, col).font = Font(name="Calibri", size=16, bold=True, color=BRAND)
    text_cell(sheet, 10, 1, "Funding by account").font = Font(name="Calibri", size=13, bold=True, color=BRAND)
    _row(sheet, 11, ["Account", f"Gross funded ({symbol})", f"Money back ({symbol})", f"Net funded ({symbol})"])
    header_row(sheet, 11, 4)
    for row, account in enumerate(data.accounts, 12):
        _row(sheet, row, [account.name, account.gross_funded, account.money_back, account.funded])
        sheet.cell(row, 3).font = Font(name="Calibri", size=11, color=RECEIVE)
    last = 12 + len(data.accounts)
    _row(sheet, last, ["Total", data.gross, data.money_back, data.net])
    _total(sheet, last, 4)
    if data.categories:
        chart_col = max(27, len(data.categories) + 6)
        category_letter = get_column_letter(chart_col)
        amount_letter = get_column_letter(chart_col + 1)
        sheet.cell(2, chart_col, "Category")
        sheet.cell(2, chart_col + 1, f"Net spending ({symbol})")
        for row, category in enumerate(data.categories, 3):
            text_cell(sheet, row, chart_col, category)
            money_cell(sheet.cell(row, chart_col + 1, data.category_totals[category]))
        chart = BarChart()
        chart.type = "bar"
        chart.title = f"Net spending by category ({data.currency})"
        chart.add_data(Reference(sheet, min_col=chart_col + 1, min_row=2, max_row=2 + len(data.categories)), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=chart_col, min_row=3, max_row=2 + len(data.categories)))
        chart.legend = None
        chart.width, chart.height = 20, min(24, max(7, 2 + len(data.categories) * 0.65))
        chart.visible_cells_only = False
        chart.y_axis.numFmt = "#,##0"
        chart.y_axis.tickLblPos = "none"
        chart.y_axis.title = f"Net spending ({symbol})"
        chart.dataLabels = DataLabelList(showVal=True, showLegendKey=False, showCatName=False, numFmt="#,##0;(#,##0);0",
            dLblPos="outEnd", dLbl=[DataLabel(idx=i, numFmt=indian_format(data.category_totals[c]))
                                    for i, c in enumerate(data.categories)])
        series = chart.series[0]
        series.graphicalProperties.solidFill = NEUTRAL
        series.graphicalProperties.line.noFill = True
        series.val.numRef.numCache = NumData(ptCount=len(data.categories), pt=[NumVal(idx=i, v=str(data.category_totals[c])) for i, c in enumerate(data.categories)])
        series.cat = AxDataSource(strRef=StrRef(f=f"'Summary'!${category_letter}$3:${category_letter}${2 + len(data.categories)}",
            strCache=StrData(ptCount=len(data.categories), pt=[StrVal(idx=i, v=c) for i, c in enumerate(data.categories)])))
        series.data_points = [DataPoint(idx=i, spPr=None) for i, c in enumerate(data.categories) if data.category_totals[c] < 0]
        for point in series.data_points:
            point.graphicalProperties.solidFill = RECEIVE
        sheet.add_chart(chart, "F10")
        sheet.column_dimensions.group(category_letter, amount_letter, hidden=True)
    pivot_start = max(last + 10, 32 + max(0, len(data.categories) - 10))
    _note(sheet, pivot_start - 6, "Person-by-category shares explain costs, not final cash balances. Headline figures and chart show the whole trip; pivot filters affect only the matrix.", 12)
    bottom = add_allocation_pivot(workbook, sheet, headers, source_rows, data, pivot_start, money_cell, header_row)
    for col, category in enumerate(data.categories, 3):
        sheet.column_dimensions[get_column_letter(col)].width = min(24, max(13, max(map(len, category.split())) + 2))
    for row in range(pivot_start + 2, bottom + 1):
        label = sheet.cell(row, 1).value
        if isinstance(label, str) and (label.endswith(" Total") or label == "Grand Total"):
            _total(sheet, row, len(data.categories) + 3)
    if source_rows:
        region = f"C{pivot_start + 2}:{get_column_letter(len(data.categories) + 3)}{bottom}"
        _number_rules(sheet, region, f"C{pivot_start + 2}")
        sheet.conditional_formatting.add(region, CellIsRule(operator="lessThan", formula=["0"], font=Font(color=RECEIVE)))
    for warning in data.warnings:
        bottom += 2
        _note(sheet, bottom, warning, 12)
    sheet.freeze_panes = "C2"
    sheet.print_area = f"A1:{get_column_letter(max(12, len(data.categories) + 3))}{bottom}"


def _members(workbook, data):
    sheet = workbook.create_sheet("Members & Families")
    symbol, has_adjustments = currency_symbol(data.currency), bool(data.adjustments)
    headers = ["Name", f"Expense share ({symbol})", f"Net funded ({symbol})", f"Before settlements ({symbol})",
               f"Settlements paid ({symbol})", f"Settlements received ({symbol})"]
    if has_adjustments:
        headers.append(f"Adjustments ({symbol})")
    headers += [f"Balance now ({symbol})", "Action"]
    _title(sheet, "Members & Families", len(headers))
    _note(sheet, 2, "Positive balance: receive. Negative balance: pay. Member positions are included in family totals; funding and settlements belong to the family account.", len(headers))
    _row(sheet, 4, headers)
    header_row(sheet, 4, len(headers))
    positions = defaultdict(list)
    for person in data.member_positions:
        positions[person.account_id].append(person)
    row = 5
    balance_col = len(headers) - 1
    for account in data.accounts:
        values = [account.name, account.share, account.funded, account.before_settlements, account.sent, account.received]
        if has_adjustments:
            values.append(account.adjustment)
        _row(sheet, row, [*values, account.balance, account.action])
        _total(sheet, row, len(headers))
        sheet.cell(row, 5).font = Font(name="Calibri", size=11, color=PAY)
        sheet.cell(row, 6).font = Font(name="Calibri", size=11, color=RECEIVE)
        start = row + 1
        row += 1
        for person in positions[account.account_id]:
            action = "Within family total" if person.balance is not None else "Historical share only"
            values = [person.name, person.share] + [None] * (balance_col - 3) + [person.balance, action]
            _row(sheet, row, values)
            sheet.cell(row, 1).alignment = Alignment(indent=2, vertical="center", wrap_text=True)
            row += 1
        if row > start:
            sheet.row_dimensions.group(start, row - 1, outline_level=1, hidden=False)
    values = ["Total", sum(a.share for a in data.accounts), sum(a.funded for a in data.accounts),
              sum(a.before_settlements for a in data.accounts), sum(a.sent for a in data.accounts), sum(a.received for a in data.accounts)]
    if has_adjustments:
        values.append(sum(a.adjustment for a in data.accounts))
    _row(sheet, row, [*values, 0, ""])
    _total(sheet, row, len(headers))
    _balance_rules(sheet, f"{get_column_letter(balance_col)}5:{get_column_letter(balance_col)}{row}")
    sheet.conditional_formatting.add(f"B5:B{row}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color=RECEIVE)))
    _number_rules(sheet, f"B5:{get_column_letter(balance_col)}{row}", "B5")
    if has_adjustments:
        row += 2
        text_cell(sheet, row, 1, "Balance reconciliation notes").font = Font(name="Calibri", size=13, bold=True, color=BRAND)
        for adjustment in data.adjustments:
            row += 1
            _note(sheet, row, f"{adjustment.account} — {adjustment.kind} ({symbol} {indian_number(adjustment.amount)}): {adjustment.explanation}", len(headers))
    sheet.sheet_properties.outlinePr.summaryBelow = False
    sheet.freeze_panes = "B5"
    _widths(sheet, [27] + [20] * (len(headers) - 2) + [23])
    sheet.print_title_rows = "1:4"
    sheet.print_area = f"A1:{get_column_letter(len(headers))}{row}"


def _payments(workbook, data):
    sheet = workbook.create_sheet("Payments")
    _title(sheet, "Payments", 7)
    _note(sheet, 2, "Suggested transfers are outstanding instructions. Recorded settlement history below contains only events already applied to balances.", 7)
    _row(sheet, 3, ["Routing", routing_label(data)])
    headers = ["Suggested payer", "Suggested receiver", f"Amount ({currency_symbol(data.currency)})", "Status"]
    _row(sheet, 5, headers)
    header_row(sheet, 5, 4)
    row = 6
    for transfer in data.transfers:
        _row(sheet, row, [transfer.payer, transfer.receiver, transfer.amount, "Suggested"])
        sheet.cell(row, 1).font = Font(name="Calibri", size=11, color=PAY)
        sheet.cell(row, 2).font = Font(name="Calibri", size=11, color=RECEIVE)
        row += 1
    if row == 6:
        text_cell(sheet, row, 1, "Everyone is settled")
        row += 1
    _row(sheet, row, ["Outstanding plan total", None, data.outstanding])
    _total(sheet, row, 4)
    row += 3
    text_cell(sheet, row, 1, "Recorded settlement history").font = Font(name="Calibri", size=13, bold=True, color=BRAND)
    row += 1
    header = row
    _row(sheet, row, ["Payer", "Receiver", f"Amount ({currency_symbol(data.currency)})", "Date & time (IST)", "Source", "Status", "Note", "Event ID"])
    header_row(sheet, row, 8)
    for event in data.events:
        row += 1
        _row(sheet, row, [event.payer, event.receiver, event.amount, event.occurred_at, event.source, event.status, event.note, event.source_id])
        sheet.cell(row, 1).font = Font(name="Calibri", size=11, color=PAY)
        sheet.cell(row, 2).font = Font(name="Calibri", size=11, color=RECEIVE)
    _table(sheet, "ReportSettlements", header, row, 8)
    for label, value in (("History total", sum(e.amount for e in data.events)),
                         ("Recorded payments total", sum(e.amount for e in data.events if e.source == "Recorded payment")),
                         ("Legacy settlements total", sum(e.amount for e in data.events if e.source == "Legacy settlement"))):
        row += 1
        _row(sheet, row, [label, None, value])
        _total(sheet, row, 7)
    _widths(sheet, [28, 28, 22, 27, 24, 24, 42, 30])
    sheet.column_dimensions["H"].hidden = True
    sheet.freeze_panes = "C6"
    sheet.print_area = f"A1:G{row}"
    sheet.print_title_rows = f"{header}:{header}"


def build_report_workbook(data: ReportData):
    workbook = Workbook()
    workbook.active.title = "Summary"
    workbook.properties.title = str(data.snapshot.trip.get("name", "Group")) + " - Report"
    workbook.properties.subject = "Trip report; report format 2"
    workbook.properties.creator = "Trip Splitter"
    headers, source_rows = _transactions(workbook, data)
    _summary(workbook, data, headers, source_rows)
    _members(workbook, data)
    _payments(workbook, data)
    workbook._sheets.sort(key=lambda sheet: ["Summary", "Members & Families", "Transactions", "Payments"].index(sheet.title))
    if data.snapshot.financial_review:
        review = workbook.create_sheet("Financial Review")
        text_cell(review, 1, 1, "Financial corrections and retained evidence")
        for row, explanation in enumerate(data.snapshot.financial_review, start=2):
            text_cell(review, row, 1, explanation)
        review.column_dimensions["A"].width = 100
    for sheet in workbook:
        finalize_layout(sheet, tuple(range(2, 12)))
    return workbook


def build_report_xlsx(data: ReportData) -> bytes:
    output = io.BytesIO()
    build_report_workbook(data).save(output)
    return output.getvalue()
